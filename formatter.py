import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill, numbers
from openpyxl.utils import get_column_letter


def formatear_excel(filename="rendimiento_acciones.xlsx"):
    wb = openpyxl.load_workbook(filename)
    ws = wb.active

    # --- 1. Ajustar ancho de columnas automáticamente ---
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # --- 2. Dar estilo al encabezado ---
    header_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    ws.row_dimensions[1].height = 28

    # --- 3. Formatear datos (filas 2 en adelante) ---
    max_row = ws.max_row

    for row in range(2, max_row + 1):
        # Precios: Formato Moneda $ #,##0.00
        ws[f"C{row}"].number_format = '"$"#,##0.00'
        ws[f"D{row}"].number_format = '"$"#,##0.00'
        ws[f"E{row}"].number_format = '"$"#,##0.00'

        # Variación %: Formato Porcentaje +0.00%;-0.00%
        var_cell = ws[f"F{row}"]
        var_cell.number_format = "+0.00%;-0.00%"

        # Color condicional según ganancia/pérdida
        val = var_cell.value
        if val is not None:
            if val > 0:
                var_cell.font = Font(color="008000", bold=True)  # Verde
            elif val < 0:
                var_cell.font = Font(color="FF0000", bold=True)  # Rojo

    # --- 4. Configurar y agregar el gráfico correctamente ---
    chart = BarChart()
    chart.type = "col"  # Columnas verticales
    chart.style = 10
    chart.title = "Rendimiento Diario (%)"
    chart.y_axis.title = "Variación (%)"
    chart.x_axis.title = "Acciones"

    # Los datos están en la Columna F (Variación %), incluyendo el título en la fila 1
    data = Reference(ws, min_col=6, min_row=1, max_row=max_row)

    # Las categorías (nombres de X) están en la Columna A (Tickers), filas 2 a max_row
    cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None  # No necesita leyenda si cada barra es un ticker en X

    chart.width = 16
    chart.height = 10

    # Colocar el gráfico al lado de la tabla
    ws.add_chart(chart, "H2")

    wb.save(filename)
